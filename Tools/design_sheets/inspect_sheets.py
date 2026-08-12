import openpyxl, json, os
base = r'D:\dOCS\test\Tools\design_sheets'
for name in ['enemies.xlsx','gods.xlsx']:
    path = os.path.join(base, name)
    wb = openpyxl.load_workbook(path, data_only=True)
    print('FILE', name)
    print('SHEETS', wb.sheetnames)
    for ws in wb.worksheets:
        print('SHEET', ws.title, 'rows', ws.max_row, 'cols', ws.max_column)
        for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
            print('\t'.join('' if v is None else str(v) for v in r))
        print('---')